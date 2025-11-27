"""
CV Tailoring Agent

Customizes user's CV/resume to match specific job requirements using LLM.
"""

import os
import json
from pathlib import Path
from typing import Dict, Any, Optional, List
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from loguru import logger

from llm import create_llm_client, BaseLLMClient


class CVTailorAgent:
    """Agent responsible for tailoring CVs to match job requirements."""

    def __init__(self, config: Dict[str, Any], llm_client: Optional[BaseLLMClient] = None):
        """
        Initialize the CV Tailor Agent.

        Args:
            config: Configuration dictionary
            llm_client: Optional pre-configured LLM client (for testing/DI)
        """
        self.config = config
        self.cv_config = config.get('cv', {})
        self.llm_config = config.get('llm', {})
        self.user_info = config.get('user', {})
        self.skills = config.get('skills', {})

        # Initialize LLM client (model-agnostic)
        if llm_client:
            self.llm_client = llm_client
        else:
            self.llm_client = create_llm_client(self.llm_config)

        # Ensure output directory exists
        self.output_dir = Path("data/cvs")
        self.output_dir.mkdir(parents=True, exist_ok=True)

        logger.info(
            f"CV Tailor Agent initialized with {self.llm_client.provider_name} "
            f"({self.llm_client.model})"
        )

    def tailor_cv(self, job: Dict[str, Any], base_cv_data: Optional[Dict] = None) -> Dict[str, Any]:
        """
        Tailor a CV for a specific job posting.

        Args:
            job: Job dictionary with description and requirements
            base_cv_data: Base CV data (if None, generates from config)

        Returns:
            Dictionary with tailored CV data and file path
        """
        logger.info(f"Tailoring CV for: {job['title']} at {job['company']}")

        try:
            # Generate base CV data if not provided
            if base_cv_data is None:
                base_cv_data = self._generate_base_cv_data()

            # Analyze job requirements
            job_analysis = self._analyze_job_requirements(job)

            # Generate tailored CV content
            tailored_cv_data = self._generate_tailored_cv_content(
                base_cv_data,
                job,
                job_analysis
            )

            # Create document
            output_path = self._create_cv_document(tailored_cv_data, job)

            result = {
                'cv_data': tailored_cv_data,
                'file_path': str(output_path),
                'job_id': job.get('url', '').split('/')[-1],
                'job_title': job['title'],
                'company': job['company'],
                'tailoring_summary': job_analysis.get('summary', ''),
                'key_skills_highlighted': job_analysis.get('key_skills', []),
                'created_at': datetime.now().isoformat()
            }

            logger.info(f"CV tailored successfully: {output_path}")

            return result

        except Exception as e:
            logger.error(f"Error tailoring CV: {e}")
            raise

    def _generate_base_cv_data(self) -> Dict[str, Any]:
        """
        Generate base CV data from configuration.

        Returns:
            Base CV data dictionary
        """
        return {
            'personal_info': {
                'name': self.user_info.get('name', ''),
                'email': self.user_info.get('email', ''),
                'phone': self.user_info.get('phone', ''),
                'linkedin': self.user_info.get('linkedin', ''),
                'github': self.user_info.get('github', ''),
            },
            'summary': '',  # Will be generated per job
            'skills': self.skills,
            'experience': self.config.get('experience', []),
            'education': self.config.get('education', []),
            'projects': self.config.get('projects', []),
        }

    def _analyze_job_requirements(self, job: Dict[str, Any]) -> Dict[str, Any]:
        """
        Analyze job requirements using LLM.

        Args:
            job: Job dictionary

        Returns:
            Analysis dictionary with key requirements and skills
        """
        logger.debug("Analyzing job requirements with LLM")

        job_description = job.get('full_description', job.get('description', ''))

        prompt = f"""Analyze this job posting and extract key information:

Job Title: {job['title']}
Company: {job['company']}
Location: {job['location']}

Job Description:
{job_description}

Please provide:
1. A list of required technical skills
2. A list of required soft skills
3. Key responsibilities
4. Experience level required
5. A brief summary of what this role is about

Format your response as JSON with these keys:
- technical_skills: list of strings
- soft_skills: list of strings
- responsibilities: list of strings
- experience_level: string
- summary: string
"""

        try:
            # Use model-agnostic LLM client
            analysis = self.llm_client.generate_json(
                prompt=prompt,
                temperature=0.3,
                max_tokens=2000
            )

            return analysis

        except Exception as e:
            logger.error(f"Error analyzing job requirements: {e}")
            # Return basic analysis
            return {
                'technical_skills': [],
                'soft_skills': [],
                'responsibilities': [],
                'experience_level': 'unknown',
                'summary': job_description[:200]
            }

    def _generate_tailored_cv_content(
        self,
        base_cv_data: Dict[str, Any],
        job: Dict[str, Any],
        job_analysis: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Generate tailored CV content using LLM.

        Args:
            base_cv_data: Base CV data
            job: Job dictionary
            job_analysis: Job requirements analysis

        Returns:
            Tailored CV data
        """
        logger.debug("Generating tailored CV content")

        user_skills_str = json.dumps(self.skills, indent=2)
        experience_str = json.dumps(base_cv_data.get('experience', []), indent=2)
        projects_str = json.dumps(base_cv_data.get('projects', []), indent=2)

        prompt = f"""You are a professional CV writer. Tailor this CV for the specific job posting.

Job Title: {job['title']}
Company: {job['company']}
Job Requirements: {json.dumps(job_analysis, indent=2)}

User's Skills:
{user_skills_str}

User's Experience:
{experience_str}

User's Projects:
{projects_str}

User's Info:
Name: {base_cv_data['personal_info']['name']}
Email: {base_cv_data['personal_info']['email']}

Task:
1. Write a compelling professional summary (3-4 sentences) that highlights relevant experience and skills for THIS specific job. Use concrete numbers and metrics where possible.
2. List 6-8 most relevant skills that match the job requirements (prioritize technical skills mentioned in the job).
3. Rewrite the "description" bullet points for each experience entry to emphasize relevance to the job. Keep them truthful but highlight matching technologies/methodologies.
4. Suggest 2-3 key achievements or experiences to emphasize (can be hypothetical but realistic based on the skills).

Format as JSON:
{{
    "professional_summary": "...",
    "key_skills": ["skill1", "skill2", ...],
    "tailored_experience": [
        {{
            "company": "Company Name",
            "title": "Job Title",
            "period": "Date Range",
            "description": ["bullet 1", "bullet 2", ...]
        }},
        ...
    ],
    "key_achievements": ["achievement1", "achievement2", ...],
    "keywords_to_emphasize": ["keyword1", "keyword2", ...]
}}

Make it specific to this role at {job['company']}. Use keywords from the job posting naturally. Do NOT use generic buzzwords. Be concise and impact-oriented.
"""

        try:
            # Use model-agnostic LLM client
            tailored_content = self.llm_client.generate_json(
                prompt=prompt,
                temperature=0.7,
                max_tokens=3000
            )

            # Merge with base CV data
            cv_data = base_cv_data.copy()
            cv_data['summary'] = tailored_content.get('professional_summary', '')
            cv_data['key_skills'] = tailored_content.get('key_skills', [])
            cv_data['key_achievements'] = tailored_content.get('key_achievements', [])
            cv_data['keywords'] = tailored_content.get('keywords_to_emphasize', [])
            
            # Update experience with tailored descriptions if available
            if tailored_content.get('tailored_experience'):
                cv_data['experience'] = tailored_content.get('tailored_experience')

            return cv_data

        except Exception as e:
            logger.error(f"Error generating tailored content: {e}")
            # Return base CV with minimal tailoring
            cv_data = base_cv_data.copy()
            cv_data['summary'] = f"Experienced professional seeking {job['title']} position at {job['company']}"
            return cv_data

    def _create_cv_document(self, cv_data: Dict[str, Any], job: Dict[str, Any]) -> Path:
        """
        Create a CV document (DOCX format).

        Args:
            cv_data: Tailored CV data
            job: Job dictionary

        Returns:
            Path to created document
        """
        logger.debug("Creating CV document")

        doc = Document()

        # Set default font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

        # Add personal info header
        personal_info = cv_data['personal_info']
        heading = doc.add_heading(personal_info['name'], level=1)
        heading.alignment = 1  # Center

        contact_info = doc.add_paragraph()
        contact_info.alignment = 1  # Center
        contact_info.add_run(
            f"{personal_info['email']} | {personal_info['phone']}\n"
            f"{personal_info.get('linkedin', '')} | {personal_info.get('github', '')}"
        )

        # Add professional summary
        doc.add_heading('Professional Summary', level=2)
        doc.add_paragraph(cv_data.get('summary', ''))

        # Add key skills
        doc.add_heading('Key Skills', level=2)
        skills_para = doc.add_paragraph()
        key_skills = cv_data.get('key_skills', [])
        skills_para.add_run(' • '.join(key_skills))

        # Add experience section
        doc.add_heading('Professional Experience', level=2)
        for exp in cv_data.get('experience', []):
            title_line = doc.add_paragraph()
            title_line.add_run(f"{exp.get('title')}").bold = True
            title_line.add_run(f" | {exp.get('company')}").italic = True
            title_line.add_run(f" | {exp.get('period')}")
            
            for bullet in exp.get('description', []):
                doc.add_paragraph(bullet, style='List Bullet')

        # Add education section
        doc.add_heading('Education', level=2)
        for edu in cv_data.get('education', []):
            edu_line = doc.add_paragraph()
            edu_line.add_run(f"{edu.get('degree')}").bold = True
            edu_line.add_run(f" | {edu.get('institution')}").italic = True
            edu_line.add_run(f" | {edu.get('period')}")

        # Add projects section
        if cv_data.get('projects'):
            doc.add_heading('Key Projects', level=2)
            for proj in cv_data.get('projects', []):
                proj_line = doc.add_paragraph()
                proj_line.add_run(f"{proj.get('name')}").bold = True
                if proj.get('technologies'):
                    techs = ", ".join(proj.get('technologies'))
                    proj_line.add_run(f" ({techs})").italic = True
                
                doc.add_paragraph(proj.get('description', ''), style='Normal')

        # Add key achievements (as extra highlights)
        if cv_data.get('key_achievements'):
            doc.add_heading('Key Achievements', level=2)
            for achievement in cv_data.get('key_achievements', []):
                doc.add_paragraph(achievement, style='List Bullet')

        # Add technical skills breakdown
        doc.add_heading('Technical Skills', level=2)
        skills = cv_data.get('skills', {})
        for category, items in skills.items():
            if isinstance(items, list) and items:
                skill_para = doc.add_paragraph()
                skill_para.add_run(f"{category.replace('_', ' ').title()}: ").bold = True
                skill_para.add_run(', '.join(items))

        # Add footer note
        doc.add_paragraph()
        footer = doc.add_paragraph()
        footer.add_run(
            f"\nTailored for: {job['title']} at {job['company']}\n"
            f"Generated on: {datetime.now().strftime('%Y-%m-%d')}"
        ).italic = True
        footer.alignment = 1  # Center

        # Save document
        job_id = job.get('url', '').split('/')[-1] or 'job'
        company_safe = ''.join(c for c in job['company'] if c.isalnum() or c in (' ', '-'))[:30]
        filename = f"CV_{company_safe}_{job_id}_{datetime.now().strftime('%Y%m%d')}.docx"
        output_path = self.output_dir / filename

        doc.save(str(output_path))

        logger.info(f"CV document created: {output_path}")

        return output_path

    def load_cv_template(self, template_path: str) -> Dict[str, Any]:
        """
        Load CV data from an existing DOCX template.

        Args:
            template_path: Path to CV template

        Returns:
            CV data dictionary
        """
        try:
            doc = Document(template_path)

            # Extract text from document
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)

            # This is a simple extraction - in production, you'd want more sophisticated parsing
            cv_data = self._generate_base_cv_data()
            cv_data['template_content'] = '\n'.join(full_text)

            return cv_data

        except Exception as e:
            logger.error(f"Error loading CV template: {e}")
            return self._generate_base_cv_data()
